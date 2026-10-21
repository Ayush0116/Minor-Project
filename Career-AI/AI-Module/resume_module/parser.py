
from io import BytesIO
from pathlib import Path
import zipfile

from docx import Document
from pypdf import PdfReader


MAX_FILE_SIZE = 5 * 1024 * 1024
MAX_EXTRACTED_CHARS = 12000


class ResumeParsingError(ValueError):
    """Raised when a resume cannot be parsed safely."""


def extract_resume_text(filename: str, content: bytes) -> str:
    if not content:
        raise ResumeParsingError("The uploaded file is empty.")

    if len(content) > MAX_FILE_SIZE:
        raise ResumeParsingError(
            "Resume file must be 5 MB or smaller."
        )

    extension = Path(filename or "").suffix.lower()

    if extension not in {".pdf", ".docx"}:
        raise ResumeParsingError(
            "Only PDF and DOCX files are supported."
        )

    try:
        if extension == ".pdf":
            if not content.startswith(b"%PDF-"):
                raise ResumeParsingError(
                    "The uploaded file is not a valid PDF."
                )

            reader = PdfReader(BytesIO(content), strict=False)

            if reader.is_encrypted:
                raise ResumeParsingError(
                    "Password-protected PDFs are not supported."
                )

            parts = []

            for page in reader.pages:
                parts.append(page.extract_text() or "")

            text = "\n".join(parts)

        else:
            stream = BytesIO(content)

            if not zipfile.is_zipfile(stream):
                raise ResumeParsingError(
                    "The uploaded file is not a valid DOCX document."
                )

            stream.seek(0)

            with zipfile.ZipFile(stream) as archive:
                if "word/document.xml" not in archive.namelist():
                    raise ResumeParsingError(
                        "The DOCX document structure is invalid."
                    )

            document = Document(BytesIO(content))
            parts = []

            for paragraph in document.paragraphs:
                if paragraph.text.strip():
                    parts.append(paragraph.text.strip())

            for table in document.tables:
                for row in table.rows:
                    cells = [
                        cell.text.strip()
                        for cell in row.cells
                        if cell.text.strip()
                    ]

                    if cells:
                        parts.append(" | ".join(cells))

            text = "\n".join(parts)

    except ResumeParsingError:
        raise
    except Exception as exc:
        raise ResumeParsingError(
            "Could not read the resume. Please check the file."
        ) from exc

    text = "\n".join(
        line.strip()
        for line in text.replace("\x00", "").splitlines()
        if line.strip()
    ).strip()

    if len(text) < 40:
        raise ResumeParsingError(
            "Very little readable text was found. "
            "Scanned PDFs may require OCR."
        )

    return text[:MAX_EXTRACTED_CHARS]
